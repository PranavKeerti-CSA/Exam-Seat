import re

with open('backend/app/routers/allocations.py', 'r') as f:
    content = f.read()

start_marker = "    allocations = db.query(models.SeatAllocation).filter("
end_marker = "        media_type=\"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\"\n    )"

start_idx = content.find(start_marker)
end_idx = content.find(end_marker) + len(end_marker)

if start_idx == -1 or end_idx == -1:
    print("Markers not found")
    exit(1)

new_logic = """    allocations = db.query(models.SeatAllocation).filter(
        models.SeatAllocation.session_id == session_id,
        models.SeatAllocation.student_id != None
    ).all()

    from openpyxl.styles import PatternFill, Font
    
    # Group by room
    room_allocs = {}
    for a in allocations:
        room_name = a.room.name if a.room else "Unknown Room"
        if room_name not in room_allocs:
            room_allocs[room_name] = []
        room_allocs[room_name].append(a)

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    with pd.ExcelWriter(temp_file.name, engine="openpyxl") as writer:
        for room_name, allocs in room_allocs.items():
            records = []
            # Sort by physical seat just in case, or by whatever order. The user said "name list".
            # Usually roll number order is preferred for a name list.
            allocs.sort(key=lambda x: x.student.roll_no if x.student else "")
            
            for a in allocs:
                stud = a.student
                sub = a.subject
                records.append({
                    "Roll Number": stud.roll_no if stud else "",
                    "Name": stud.name if stud else "",
                    "Class": stud.grade if stud else "",
                    "Subject": sub.name if sub else "",
                    "_color": sub.color if sub else "#FFFFFF"
                })
            
            df = pd.DataFrame(records)
            colors = df.pop("_color").tolist() if "_color" in df.columns else []
            
            # Write to excel
            safe_room_name = str(room_name).replace("/", "_")[:31]
            df.to_excel(writer, sheet_name=safe_room_name, startrow=2, index=False)
            worksheet = writer.sheets[safe_room_name]
            
            # Add Date at top
            worksheet.cell(row=1, column=1, value=f"Exam Date: {session.date}")
            worksheet.cell(row=1, column=1).font = Font(bold=True, size=14)
            
            # Apply color coding to rows
            for r_idx, row in enumerate(worksheet.iter_rows(min_row=4, max_row=3 + len(df), min_col=1, max_col=4), start=0):
                hex_color = colors[r_idx].lstrip('#') if r_idx < len(colors) else "FFFFFF"
                if len(hex_color) == 6:
                    # lighten color slightly if it's too dark so text is readable
                    # actually, just use the raw color.
                    fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
                    for cell in row:
                        cell.fill = fill

    clean_name = session.name.replace(" ", "_")
    return FileResponse(
        path=temp_file.name,
        filename=f"{clean_name}_Classwise_Lists.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )"""

new_content = content[:start_idx] + new_logic + content[end_idx:]
with open('backend/app/routers/allocations.py', 'w') as f:
    f.write(new_content)
print("Export patched")
